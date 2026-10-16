"""Text extraction and chunking: the bytes of a supported file become chunks with source locations.

Pure CPU work without I/O. Extractors are generators, so limits stop oversized documents early.
"""

from __future__ import annotations

import io
import logging
import re
import zipfile
from collections.abc import Callable, Iterable, Iterator
from dataclasses import dataclass, field
from typing import TypeVar

from docx import Document as WordDocument
from docx.table import Table as WordTable
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.shapes.group import GroupShape
from pypdf import PasswordType, PdfReader

logger = logging.getLogger("rag.extraction")

MAX_CHUNK_CHARS = 1200
# Bounds memory, embedding work and rows per document (about 24 MB of text).
MAX_CHUNKS = 20_000
# Zip bomb guards on the declared uncompressed size of OOXML files. DOCX and PPTX parsers build
# the whole XML tree (roughly 20x its size in memory); XLSX is read as a stream.
MAX_UNPACKED_TREE_BYTES = 32 * 1024 * 1024
MAX_UNPACKED_STREAM_BYTES = 128 * 1024 * 1024

MARKDOWN_HEADING = re.compile(r"^ {0,3}#{1,6}[ \t]+(.+?)(?:[ \t]+#+)?[ \t]*$")
MARKDOWN_FENCE = re.compile(r"^ {0,3}(?:```|~~~)")
# Only real line breaks: str.splitlines() also splits on form feeds and  , shifting line numbers.
LINE_BREAK = re.compile(r"\r\n|\r|\n")
SEARCHABLE = re.compile(r"\w")

T = TypeVar("T")


class ExtractionError(Exception):
    """The file cannot be indexed; the message is safe to return to the client."""


@dataclass(frozen=True, slots=True)
class Units:
    """A run of numbered units of one kind: line 3, paragraphs 2-5, page 7."""

    unit: str
    first: int
    last: int

    def merge(self, other: Units) -> Units:
        return Units(self.unit, self.first, other.last)

    def __str__(self) -> str:
        if self.first == self.last:
            return f"{self.unit} {self.first}"
        return f"{self.unit}s {self.first}-{self.last}"


@dataclass(frozen=True, slots=True)
class Cells:
    """The spreadsheet range covering every packed row."""

    min_row: int
    min_column: int
    max_row: int
    max_column: int

    def merge(self, other: Cells) -> Cells:
        return Cells(
            min(self.min_row, other.min_row),
            min(self.min_column, other.min_column),
            max(self.max_row, other.max_row),
            max(self.max_column, other.max_column),
        )

    def __str__(self) -> str:
        start = f"{get_column_letter(self.min_column)}{self.min_row}"
        end = f"{get_column_letter(self.max_column)}{self.max_row}"
        return f"cells {start}:{end}"


@dataclass(frozen=True, slots=True)
class Block:
    """One extracted unit. Consecutive packable blocks of the same scope share a chunk."""

    text: str
    span: Units | Cells
    scope: str = ""  # 'section "Pricing"', 'sheet Budget', 'section "Limits", table 1'
    packable: bool = True  # pages and slides stay separate chunks


@dataclass(frozen=True, slots=True)
class Chunk:
    text: str
    location: str


@dataclass(slots=True)
class Extraction:
    chunks: list[Chunk]
    # Parts that could not be read and were skipped, so the rest of the document is still indexed.
    warnings: list[str] = field(default_factory=list)


def join_scope(*parts: str) -> str:
    return ", ".join(part for part in parts if part)


def join_cells(texts: Iterable[str]) -> str:
    return " | ".join(text for text in texts if text.strip())


def iter_lines(text: str) -> Iterator[str]:
    start = 0
    for line_break in LINE_BREAK.finditer(text):
        yield text[start : line_break.start()]
        start = line_break.end()
    yield text[start:]


def split_text(text: str) -> Iterator[str]:
    """Cut text longer than MAX_CHUNK_CHARS at whitespace, without copying the remainder each time."""
    start, end = 0, len(text)
    while start < end:
        stop = min(start + MAX_CHUNK_CHARS, end)
        if stop < end:
            cut = max(text.rfind(" ", start, stop), text.rfind("\n", start, stop))
            if cut > start + MAX_CHUNK_CHARS // 2:
                stop = cut
        piece = text[start:stop].strip()
        if piece:
            yield piece
        start = stop


class ChunkBuilder:
    """Packs consecutive small blocks of one scope into chunks of up to MAX_CHUNK_CHARS.

    Packing keeps the chunk count proportional to the amount of text rather than to the number
    of lines or rows: a 0.4 MB file of short lines used to become 200 000 chunks and 3 GB of RAM.
    """

    def __init__(self) -> None:
        self.chunks: list[Chunk] = []
        self._seen: set[str] = set()
        self._texts: list[str] = []
        self._span: Units | Cells | None = None
        self._scope = ""
        self._size = 0

    def add(self, block: Block) -> None:
        text = block.text.strip()
        if not SEARCHABLE.search(text):
            return  # blank or punctuation only: its embedding carries no meaning
        if not block.packable or len(text) > MAX_CHUNK_CHARS:
            self.flush()
            location = self._location(block.scope, block.span)
            for number, piece in enumerate(split_text(text), start=1):
                self._emit(piece, location if number == 1 else f"{location}, part {number}")
            return
        if self._span is not None and (block.scope != self._scope or self._size + len(text) + 1 > MAX_CHUNK_CHARS):
            self.flush()
        if self._span is None:
            self._span, self._scope = block.span, block.scope
        else:
            self._span = self._span.merge(block.span)
        self._texts.append(text)
        self._size += len(text) + 1

    def flush(self) -> None:
        if self._span is not None:
            self._emit("\n".join(self._texts), self._location(self._scope, self._span))
        self._texts, self._span, self._scope, self._size = [], None, "", 0

    def _emit(self, text: str, location: str) -> None:
        # A repeated chunk is kept once, at its first location. Thousands of identical vectors
        # degrade the HNSW graph until searches stop reaching other documents at all.
        if text in self._seen:
            return
        self._seen.add(text)
        if len(self.chunks) >= MAX_CHUNKS:
            raise ExtractionError("Document is too large to index")
        self.chunks.append(Chunk(text, location))

    @staticmethod
    def _location(scope: str, span: Units | Cells) -> str:
        return join_scope(scope, str(span))


def decode_utf8(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")  # -sig drops the BOM that Windows editors put first
    except UnicodeDecodeError as error:
        raise ExtractionError("File must be valid UTF-8") from error


def check_unpacked_size(content: bytes, limit: int) -> None:
    """Reject an OOXML zip bomb before a parser inflates it; zipfile never reads past declared sizes."""
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        unpacked = sum(entry.file_size for entry in archive.infolist())
    if unpacked > limit:
        raise ExtractionError("Document is too large to process")


def tolerant(parts: Iterable[T], unit: str, warnings: list[str], read: Callable[[int, T], Iterable[Block]]) -> Iterator[Block]:
    """Yield the blocks of every readable part; a broken page, slide or sheet is skipped and reported."""
    for number, part in enumerate(parts, start=1):
        try:
            yield from read(number, part)
        except ExtractionError:
            raise
        except Exception as error:
            logger.warning("Skipping unreadable %s %d: %s", unit, number, error)
            warnings.append(f"{unit} {number} could not be read and was skipped")


def text_blocks(content: bytes, _warnings: list[str]) -> Iterator[Block]:
    for number, line in enumerate(iter_lines(decode_utf8(content)), start=1):
        yield Block(line, Units("line", number, number))


def markdown_blocks(content: bytes, _warnings: list[str]) -> Iterator[Block]:
    section = ""
    in_code = False
    for number, line in enumerate(iter_lines(decode_utf8(content)), start=1):
        if MARKDOWN_FENCE.match(line):
            in_code = not in_code
        elif not in_code and (heading := MARKDOWN_HEADING.match(line)):
            section = f'section "{heading.group(1)}"'
        yield Block(line, Units("line", number, number), section)


def docx_row_text(row) -> str:
    # row.cells repeats a horizontally merged cell once per grid column it spans.
    seen: set[int] = set()
    texts = []
    for cell in row.cells:
        if id(cell._tc) not in seen:
            seen.add(id(cell._tc))
            texts.append(cell.text)
    return join_cells(texts)


def docx_blocks(content: bytes, _warnings: list[str]) -> Iterator[Block]:
    check_unpacked_size(content, MAX_UNPACKED_TREE_BYTES)
    document = WordDocument(io.BytesIO(content))
    # paragraph.style re-resolves the style in styles.xml on every call (~200 us; minutes for a
    # document of tiny paragraphs), so style ids are mapped to names once per document.
    style_names = {style.style_id: style.name for style in document.styles}
    # Paragraphs and tables in document order, so a table gets the section it appears in.
    section = ""
    paragraphs = tables = 0
    for item in document.iter_inner_content():
        if isinstance(item, WordTable):
            tables += 1
            scope = join_scope(section, f"table {tables}")
            for number, row in enumerate(item.rows, start=1):
                yield Block(docx_row_text(row), Units("row", number, number), scope)
        else:
            paragraphs += 1
            style = style_names.get(item._p.style) or ""
            if style.startswith("Heading") and item.text.strip():
                section = f'section "{item.text.strip()}"'
            yield Block(item.text, Units("paragraph", paragraphs, paragraphs), section)


def pdf_blocks(content: bytes, warnings: list[str]) -> Iterator[Block]:
    reader = PdfReader(io.BytesIO(content))
    # PdfReader already tries the empty password, which opens permission-only encrypted PDFs.
    if reader.is_encrypted and reader.decrypt("") == PasswordType.NOT_DECRYPTED:
        raise ExtractionError("File is password-protected")
    yield from tolerant(
        reader.pages,
        "page",
        warnings,
        lambda number, page: [Block(page.extract_text() or "", Units("page", number, number), packable=False)],
    )


def sheet_blocks(_number: int, sheet) -> Iterator[Block]:
    scope = f"sheet {sheet.title}"
    for row in sheet.iter_rows():
        filled = [cell for cell in row if cell.value is not None]  # read-only rows pad gaps with EmptyCell
        if filled:
            text = " | ".join(f"{cell.coordinate}: {cell.value}" for cell in filled)
            yield Block(text, Cells(filled[0].row, filled[0].column, filled[-1].row, filled[-1].column), scope)


def xlsx_blocks(content: bytes, warnings: list[str]) -> Iterator[Block]:
    check_unpacked_size(content, MAX_UNPACKED_STREAM_BYTES)
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        yield from tolerant(workbook.worksheets, "sheet", warnings, sheet_blocks)
    finally:
        workbook.close()


def pptx_shape_texts(shapes) -> list[str]:
    texts: list[str] = []
    for shape in shapes:
        if isinstance(shape, GroupShape):
            texts.extend(pptx_shape_texts(shape.shapes))
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                values = join_cells(cell.text for cell in row.cells)
                if values:
                    texts.append(values)
        elif getattr(shape, "has_text_frame", False) and shape.text.strip():
            texts.append(shape.text)
    return texts


def pptx_blocks(content: bytes, warnings: list[str]) -> Iterator[Block]:
    check_unpacked_size(content, MAX_UNPACKED_TREE_BYTES)
    presentation = Presentation(io.BytesIO(content))
    yield from tolerant(
        presentation.slides,
        "slide",
        warnings,
        lambda number, slide: [Block("\n".join(pptx_shape_texts(slide.shapes)), Units("slide", number, number), packable=False)],
    )


# The single source of truth for supported formats.
EXTRACTORS: dict[str, Callable[[bytes, list[str]], Iterator[Block]]] = {
    ".txt": text_blocks,
    ".md": markdown_blocks,
    ".markdown": markdown_blocks,
    ".docx": docx_blocks,
    ".pdf": pdf_blocks,
    ".xlsx": xlsx_blocks,
    ".pptx": pptx_blocks,
}


def extract(suffix: str, content: bytes) -> Extraction:
    """Chunk a document of a supported format; raises ExtractionError when nothing can be indexed."""
    warnings: list[str] = []
    builder = ChunkBuilder()
    try:
        for block in EXTRACTORS[suffix](content, warnings):
            builder.add(block)
        builder.flush()
    except ExtractionError:
        raise
    except Exception as error:
        raise ExtractionError("File is damaged or unreadable") from error
    if not builder.chunks:
        if warnings:
            raise ExtractionError("File is damaged or unreadable")
        raise ExtractionError("File contains no extractable text")
    return Extraction(builder.chunks, warnings)
