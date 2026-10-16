from __future__ import annotations

import io
import zipfile

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas


def make_document(extension: str, text: str) -> bytes:
    if extension == "txt":
        return text.encode()
    if extension == "md":
        return f"# Test document\n\n{text}\n".encode()
    if extension == "docx":
        document = Document()
        document.add_heading("Test document", level=1)
        document.add_paragraph(text)
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()
    if extension == "pdf":
        output = io.BytesIO()
        page = canvas.Canvas(output)
        page.drawString(72, 720, text)
        page.save()
        return output.getvalue()
    if extension == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Evidence"
        sheet["A1"] = "fact"
        sheet["B1"] = text
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()
    if extension == "pptx":
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
        box.text = text
        output = io.BytesIO()
        presentation.save(output)
        return output.getvalue()
    raise ValueError(extension)


def make_markdown_with_sections(preamble_text: str, section_text: str) -> bytes:
    lines = [
        f"Preamble {preamble_text}",                # line 1: before any heading
        "# Overview",                                # line 2
        "General introduction.",                     # line 3
        "## Pricing",                                # line 4
        "```bash",                                   # line 5
        "# not a heading, just a shell comment",     # line 6
        "```",                                       # line 7
        f"The {section_text} plan costs 10 credits.",  # line 8: section "Pricing"
    ]
    return "\n".join(lines).encode()


def make_docx_with_sections(paragraph_text: str, table_text: str, appendix_text: str) -> bytes:
    document = Document()
    document.add_paragraph("Front matter")             # paragraph 1: before any heading
    document.add_heading("Limits", level=1)            # paragraph 2
    document.add_paragraph(paragraph_text)             # paragraph 3: section "Limits"
    table = document.add_table(rows=1, cols=2)         # table 1: section "Limits"
    table.cell(0, 0).text = "limit"
    table.cell(0, 1).text = table_text
    document.add_heading("Appendix", level=2)          # paragraph 4
    document.add_paragraph(appendix_text)              # paragraph 5: section "Appendix"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def make_pdf_without_text_layer() -> bytes:
    """A page with graphics only, like a scan that was never OCR'd."""
    output = io.BytesIO()
    page = canvas.Canvas(output)
    page.rect(72, 600, 300, 120, fill=1)
    page.line(72, 560, 372, 560)
    page.save()
    return output.getvalue()


def make_sparse_xlsx(text: str) -> bytes:
    """Sheet whose rows have empty cells before and after the values, like real spreadsheets."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sparse"
    sheet["A1"] = "name"
    sheet["B1"] = "value"
    sheet["C1"] = "note"
    sheet["A2"] = "only first column"
    sheet["B3"] = text
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def make_pptx_with_table_and_group(table_text: str, group_text: str) -> bytes:
    """Slide 1 holds text only in a table, slide 2 only inside a grouped shape."""
    presentation = Presentation()
    table_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = table_slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(8), Inches(1)).table
    table.cell(0, 0).text = "fact"
    table.cell(0, 1).text = table_text
    group_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = group_slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text = group_text
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def rewrite_zip_entry(archive: bytes, name: str, transform) -> bytes:
    """Copy an OOXML package, replacing one entry with transform(original bytes)."""
    source = zipfile.ZipFile(io.BytesIO(archive))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            target.writestr(entry, transform(data) if entry.filename == name else data)
    return output.getvalue()


def make_docx_with_merged_cells(text: str) -> bytes:
    """A table row whose first cell spans two grid columns."""
    document = Document()
    table = document.add_table(rows=1, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = text
    table.cell(0, 2).text = "other"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def make_docx_zip_bomb(unpacked_bytes: int) -> bytes:
    """A valid DOCX whose document.xml inflates to unpacked_bytes of one-letter paragraphs."""
    paragraph = b"<w:p><w:r><w:t>a</w:t></w:r></w:p>"
    body = paragraph * (unpacked_bytes // len(paragraph) + 1)
    return rewrite_zip_entry(
        make_document("docx", "seed"), "word/document.xml", lambda xml: xml.replace(b"<w:body>", b"<w:body>" + body, 1)
    )


def make_encrypted_pdf(text: str, user_password: str) -> bytes:
    """AES-128 encrypted PDF; an empty user password restricts permissions only."""
    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(make_document("pdf", text))))
    writer.encrypt(user_password=user_password, owner_password="owner-secret", algorithm="AES-128")
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def make_xlsx_with_broken_sheet(text: str) -> bytes:
    """Sheet 1 is intact; the XML of sheet 2 is cut in half, as in a partially corrupted file."""
    workbook = Workbook()
    workbook.active.title = "Intact"
    workbook.active["A1"] = text
    broken = workbook.create_sheet("Broken")
    for row in range(1, 200):
        broken[f"A{row}"] = f"value {row}"
    output = io.BytesIO()
    workbook.save(output)
    return rewrite_zip_entry(output.getvalue(), "xl/worksheets/sheet2.xml", lambda xml: xml[: len(xml) // 2])
